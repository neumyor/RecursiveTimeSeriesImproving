from __future__ import annotations

from pathlib import Path

from docx import Document


def extract_docx(path: Path) -> str:
    document = Document(str(path))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        chunks.append(f"\n[Table {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))
    return "\n\n".join(chunks)
